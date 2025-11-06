# worklog_core_multi_user.py
# -*- coding: utf-8 -*-
"""
Phiên bản multi-user: đọc nhiều CSV (mỗi nhân viên), gộp & xuất báo cáo.
ENHANCED UI VERSION - Giao diện được nâng cấp với thiết kế hiện đại

Tính năng chính
- Chuẩn hoá tên dự án theo "khóa dự án" (_normalize_project_key).
- Project Matrix: mỗi sheet = 1 dự án, luôn đủ TẤT CẢ nhân viên và TẤT CẢ ngày.
  Cột 'Tổng' nằm NGAY SAU cột tên (trước các cột ngày).
- Dashboard: tổng tăng ca tính đúng (cộng theo từng nhân viên), có mục Cảnh báo & Lưu ý cho toàn bộ nhân viên.
"""

from pathlib import Path
from datetime import datetime
from typing import List, Dict, Optional
import pandas as pd

# ==== Import từ core gốc ====
from worklog_core import (
    _parse_date_safely,
    parse_csv_data,
    extract_work_records,
    analyze_overtime,
    _unique_sheet_name,
    _pick_font_path,
    _normalize_project_key,  # dùng để canonicalize dự án
)

# ==== Canonicalize tên dự án (chỉ xét chữ) ====
_PROJECT_CANON: Dict[str, str] = {}  # key chuẩn -> tên hiển thị đầu tiên


def _canon_project_display(name: Optional[str]) -> str:
    """Gộp các biến thể dự án về cùng 1 tên hiển thị, khóa theo _normalize_project_key."""
    base = "" if name is None else str(name).strip()
    key = _normalize_project_key(base)
    if key not in _PROJECT_CANON:
        _PROJECT_CANON[key] = base
    return _PROJECT_CANON[key]


# ======================================================================
#                           MULTI-USER PROCESSING
# ======================================================================


def process_multiple_csvs(
    file_paths: List[str], DATE_FORMAT_IN_CSV: str = "day_first"
) -> Dict[str, object]:
    """
    Đọc toàn bộ CSV, chuẩn hoá tên dự án, gom dữ liệu.

    Returns:
        {
            'all_records': List[dict(name,date,project,hours)],
            'all_employees': Dict[str, List[record]],
            'dataframes': Dict[str, pd.DataFrame],
            'all_dates': List[str: 'YYYY-MM-DD'] (unique, sorted)
        }
    """
    all_records: List[dict] = []
    all_employees: Dict[str, List[dict]] = {}
    dataframes: Dict[str, pd.DataFrame] = {}
    all_dates_set: set = set()

    for fp in (file_paths or []):
        p = Path(str(fp))
        if not p.exists():
            continue

        try:
            df, employee_name = parse_csv_data(str(p), DATE_FORMAT_IN_CSV)
            records = extract_work_records(df, employee_name, DATE_FORMAT_IN_CSV)

            # Chuẩn hoá tên dự án
            for r in records:
                r["project"] = _canon_project_display(r.get("project", ""))

            if records:
                all_records.extend(records)
                all_employees[employee_name] = records
                dataframes[employee_name] = df

                # Thu thập ngày từ cột 日付 để làm cột ma trận
                if "日付" in df.columns:
                    for val in df["日付"]:
                        if pd.isna(val) or str(val).strip() == "":
                            continue
                        d = _parse_date_safely(val, DATE_FORMAT_IN_CSV)
                        if pd.notna(d):
                            all_dates_set.add(d.strftime("%Y-%m-%d"))

            print(f"✓ Đã xử lý: {employee_name} ({len(records)} records)")
        except Exception as e:
            print(f"⚠ Lỗi xử lý {p.name}: {e}")

    all_dates = sorted(all_dates_set)
    return {
        "all_records": all_records,
        "all_employees": all_employees,
        "dataframes": dataframes,
        "all_dates": all_dates,
    }


def calculate_multi_user_statistics(
    all_records: List[dict],
    all_employees: Dict[str, List[dict]],
    STANDARD_HOURS: float = 8.0,
    OVERTIME_THRESHOLD: float = 8.0,
) -> Dict[str, object]:
    """Tính thống kê tổng hợp nhiều nhân viên (overtime tổng = tổng overtime của từng nhân viên)."""
    if not all_records:
        return {}

    df_all = pd.DataFrame(all_records)
    df_all["date_obj"] = pd.to_datetime(df_all["date"])
    df_all["project"] = df_all["project"].apply(_canon_project_display)

    total_hours = float(df_all["hours"].sum())
    total_work_days = int(df_all["date"].nunique())
    total_employees = len(all_employees)
    avg_hours_per_day = float(total_hours / total_work_days) if total_work_days else 0.0

    # Theo nhân viên & tổng overtime đúng chuẩn
    employee_stats = []
    warnings_all: List[str] = []
    total_overtime_sum = 0.0

    for name, recs in all_employees.items():
        emp_df = pd.DataFrame(recs)
        emp_hours = float(emp_df["hours"].sum())
        emp_days = int(emp_df["date"].nunique())
        emp_avg = float(emp_hours / emp_days) if emp_days else 0.0

        emp_ot = analyze_overtime(recs, STANDARD_HOURS, OVERTIME_THRESHOLD)
        total_overtime_sum += float(emp_ot.get("total_overtime", 0.0))
        for w in emp_ot.get("warnings", []) or []:
            warnings_all.append(f"{name}: {w}")

        employee_stats.append(
            {
                "name": name,
                "total_hours": emp_hours,
                "work_days": emp_days,
                "avg_hours_per_day": emp_avg,
                "total_overtime": emp_ot.get("total_overtime", 0.0),
                "percentage": (emp_hours / total_hours * 100) if total_hours > 0 else 0.0,
            }
        )

    df_emp_stats = pd.DataFrame(employee_stats).sort_values("total_hours", ascending=False)

    # Theo dự án
    project_stats = (
        df_all.groupby("project", as_index=False)
        .agg(
            total_hours=("hours", "sum"),
            days_worked=("date", "nunique"),
            employees=("name", lambda x: len(set(x))),
        )
        .sort_values("total_hours", ascending=False)
    )
    project_stats["percentage"] = (
        project_stats["total_hours"] / total_hours * 100 if total_hours > 0 else 0.0
    )

    # Theo tuần
    df_all["week"] = df_all["date_obj"].dt.isocalendar().week
    weekly_stats = (
        df_all.groupby("week", as_index=False)["hours"].sum().rename(columns={"hours": "total_hours"})
    )

    overtime_info = {
        "overtime_records": [],  # có thể bổ sung khi cần
        "warnings": warnings_all,
        "total_overtime": total_overtime_sum,
    }

    return {
        "total_hours": total_hours,
        "total_work_days": total_work_days,
        "total_employees": total_employees,
        "avg_hours_per_day": avg_hours_per_day,
        "employee_stats": df_emp_stats,
        "project_stats": project_stats,
        "weekly_stats": weekly_stats,
        "overtime_info": overtime_info,
        "date_range": {
            "start": df_all["date_obj"].min().strftime("%Y-%m-%d"),
            "end": df_all["date_obj"].max().strftime("%Y-%m-%d"),
        },
    }


# ======================================================================
#                               OUTPUTS
# ======================================================================


def create_multi_user_matrix_excel(
    all_employees: Dict[str, List[dict]],
    all_dataframes: Dict[str, pd.DataFrame],
    OUT_MATRIX_XLSX: str,
    DATE_FORMAT_IN_CSV: str = "day_first",
) -> None:
    """
    Tạo project_matrix.xlsx (multi-user).
    - Mỗi sheet = 1 dự án
    - Luôn in đủ TẤT CẢ nhân viên từ input (kể cả không có giờ ở dự án)
    - Cột 'Tổng' đặt NGAY SAU cột Tên, trước các cột ngày
    - Cột ngày đầy đủ theo tất cả CSV (dd/mm, sắp theo thời gian)
    """
    if not all_employees:
        print("⚠️ Không có dữ liệu để tạo ma trận.")
        return

    # Gộp records
    all_records: List[dict] = []
    for recs in all_employees.values():
        if recs:
            all_records.extend(recs)

    if not all_records:
        print("⚠️ Không có bản ghi hợp lệ.")
        return

    # Danh sách nhân viên theo thứ tự input
    all_names = list(all_employees.keys())

    # Lấy tất cả ngày từ các CSV để làm cột dd/mm theo thời gian
    all_dates_dict: Dict[str, pd.Timestamp] = {}
    for df in all_dataframes.values():
        if df is None or "日付" not in df.columns:
            continue
        for date_val in df["日付"]:
            if pd.isna(date_val) or str(date_val).strip() == "":
                continue
            d = _parse_date_safely(date_val, DATE_FORMAT_IN_CSV)
            if pd.notna(d):
                all_dates_dict[d.strftime("%d/%m")] = d
    all_dates_in_csv = [k for k, v in sorted(all_dates_dict.items(), key=lambda x: x[1])]

    # DataFrame hợp nhất
    df_all = pd.DataFrame(all_records)
    df_all["project"] = df_all["project"].apply(_canon_project_display)
    df_all["date_obj"] = pd.to_datetime(df_all["date"])
    df_all["date_formatted"] = df_all["date_obj"].dt.strftime("%d/%m")

    projects = list(pd.unique(df_all["project"]))
    used_sheet_names = set()

    try:
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        with pd.ExcelWriter(OUT_MATRIX_XLSX, engine="openpyxl") as writer:
            for project in projects:
                sub = df_all[df_all["project"] == project].copy()
                pivot = sub.pivot_table(
                    index="name",
                    columns="date_formatted",
                    values="hours",
                    aggfunc="sum",
                    fill_value=0,
                )

                # Đảm bảo đủ cột ngày theo toàn bộ CSV
                pivot = pivot.reindex(columns=all_dates_in_csv, fill_value=0)
                # Đảm bảo đủ hàng nhân viên
                pivot = pivot.reindex(index=all_names, fill_value=0)

                # === CỘT 'Tổng' đặt NGAY SAU cột Tên ===
                row_totals = pivot.sum(axis=1, numeric_only=True)
                pivot.insert(0, "Tổng", row_totals)  # chèn làm cột đầu
                pivot = pivot.reindex(columns=["Tổng"] + all_dates_in_csv)

                # Hàng 'TỔNG' (tổng theo cột, gồm cả cột 'Tổng')
                pivot.loc["TỔNG"] = pivot.sum(numeric_only=True)

                sheet_name = _unique_sheet_name(project, used_sheet_names)
                
                # ... sau khi đã chèn cột 'Tổng' và hàng 'TỔNG'
# pivot.insert(0, "Tổng", row_totals)
# pivot = pivot.reindex(columns=["Tổng"] + all_dates_in_csv)
# pivot.loc["TỔNG"] = pivot.sum(numeric_only=True)

# --> THÊM ĐOẠN NÀY:
                def _blank_zeros(x):
    # Nếu là số và bằng 0.0 thì để trống
                    try:
                        return "" if (isinstance(x, (int, float)) and float(x) == 0.0) else x
                    except Exception:
                        return x

                pivot = pivot.applymap(_blank_zeros)

                pivot.to_excel(writer, sheet_name=sheet_name)

                # Styling
                ws = writer.sheets[sheet_name]
                header_fill = PatternFill(start_color="667EEA", end_color="667EEA", fill_type="solid")
                header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
                total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
                total_font = Font(name="Segoe UI", size=10, bold=True, color="334155")
                data_font = Font(name="Segoe UI", size=10, color="475569")
                name_font = Font(name="Segoe UI", size=10, bold=True, color="334155")
                center = Alignment(horizontal="center", vertical="center")
                left = Alignment(horizontal="left", vertical="center")
                thin = Border(
                    left=Side(style="thin", color="E2E8F0"),
                    right=Side(style="thin", color="E2E8F0"),
                    top=Side(style="thin", color="E2E8F0"),
                    bottom=Side(style="thin", color="E2E8F0"),
                )

                # Header row
                for c in range(1, ws.max_column + 1):
                    cell = ws.cell(row=1, column=c)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = center
                    cell.border = thin

                # Body
                for r in range(2, ws.max_row + 1):
                    name_cell = ws.cell(row=r, column=1)
                    if str(name_cell.value) == "TỔNG":
                        name_cell.fill = total_fill
                        name_cell.font = total_font
                    else:
                        name_cell.font = name_font
                    name_cell.alignment = left
                    name_cell.border = thin

                    for c in range(2, ws.max_column + 1):
                        cell = ws.cell(row=r, column=c)
                        hdr = str(ws.cell(row=1, column=c).value)
                        if str(name_cell.value) == "TỔNG" or hdr == "Tổng":
                            cell.fill = total_fill
                            cell.font = total_font
                        else:
                            cell.font = data_font
                        cell.alignment = center
                        cell.border = thin
                        if isinstance(cell.value, (int, float)):
                            cell.number_format = "0.0"

                # Widths + freeze
                ws.column_dimensions["A"].width = 25  # cột tên
                for c in range(2, ws.max_column + 1):
                    col_letter = get_column_letter(c)
                    if ws.cell(row=1, column=c).value == "Tổng":
                        ws.column_dimensions[col_letter].width = 12
                    else:
                        ws.column_dimensions[col_letter].width = 10
                ws.freeze_panes = "B2"

                print(f"  ✓ Sheet: {sheet_name}")

        print(f"✅ Đã tạo file ma trận tổng hợp: {OUT_MATRIX_XLSX}")
    except Exception as e:
        print(f"⚠️ Lỗi khi tạo file ma trận: {e}")


def create_multi_user_dashboard(
    stats: Dict[str, object],
    OUT_DASHBOARD_HTML: str,
    all_records: Optional[List[dict]] = None,
    all_dataframes: Optional[Dict[str, pd.DataFrame]] = None,
    DATE_FORMAT_IN_CSV: str = "day_first",
) -> None:
    """Dashboard HTML tổng hợp nhiều user (Chart.js) + Cảnh báo & Lưu ý - ENHANCED UI."""
    import json

    proj = stats["project_stats"]
    emp = stats["employee_stats"]
    week = stats["weekly_stats"]
    ot = stats["overtime_info"]

    proj_labels = proj["project"].tolist()
    proj_hours = proj["total_hours"].astype(float).tolist()

    emp_labels = emp["name"].tolist()
    emp_hours = emp["total_hours"].astype(float).tolist()

    weekly_labels = [f"Tuần {int(w)}" for w in week["week"].tolist()]
    weekly_hours = week["total_hours"].astype(float).tolist()

    warnings_list = (ot.get("warnings") if isinstance(ot, dict) else None) or []
    warnings_html = ""
    if warnings_list:
        items = "".join([f'<div class="warning-item"><span class="warning-icon">⚠️</span><span>{w}</span></div>' for w in warnings_list])
        warnings_html = f"""
        <div class="card warning-card">
          <div class="card-header">
            <h3><span class="icon">⚠️</span>Cảnh báo & Lưu ý</h3>
          </div>
          <div class="warnings-container">
            {items}
          </div>
        </div>"""

    palette = [
        "#6366F1", "#22C55E", "#F59E0B", "#EF4444", "#14B8A6", "#A855F7",
        "#3B82F6", "#84CC16", "#F97316", "#06B6D4", "#EC4899", "#10B981"
    ]
    proj_colors = [palette[i % len(palette)] for i in range(len(proj_labels))]
    emp_colors = [palette[i % len(palette)] for i in range(len(emp_labels))]
    weekly_colors = [palette[i % len(palette)] for i in range(len(weekly_labels))]

    html = f"""<!DOCTYPE html>
<html lang="vi">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Báo Cáo Tổng Hợp - Multi User</title>
<script src="https://cdnjs.cloudflare.com/ajax/libs/Chart.js/3.9.1/chart.min.js"></script>
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap" rel="stylesheet">
<style>
* {{
  margin: 0;
  padding: 0;
  box-sizing: border-box;
}}

body {{
  font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
  background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
  padding: 40px 24px;
  min-height: 100vh;
  color: #1e293b;
  line-height: 1.6;
}}

.container {{
  max-width: 1600px;
  margin: 0 auto;
}}

/* ===== HEADER ===== */
.header {{
  background: #ffffff;
  border-radius: 24px;
  box-shadow: 0 20px 60px rgba(0, 0, 0, 0.12);
  padding: 48px;
  margin-bottom: 40px;
  position: relative;
  overflow: hidden;
}}

.header::before {{
  content: '';
  position: absolute;
  top: 0;
  left: 0;
  right: 0;
  height: 6px;
  background: linear-gradient(90deg, #667eea 0%, #764ba2 100%);
}}

.header-title {{
  font-size: 42px;
  font-weight: 800;
  background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
  -webkit-background-clip: text;
  -webkit-text-fill-color: transparent;
  background-clip: text;
  margin-bottom: 24px;
  letter-spacing: -0.5px;
}}

.header-meta {{
  display: flex;
  flex-wrap: wrap;
  gap: 16px;
  align-items: center;
}}

.badge {{
  display: inline-flex;
  align-items: center;
  gap: 8px;
  padding: 12px 24px;
  border-radius: 16px;
  background: linear-gradient(135deg, #f8fafc 0%, #e2e8f0 100%);
  color: #475569;
  font-weight: 600;
  font-size: 15px;
  border: 2px solid #e2e8f0;
  transition: all 0.3s ease;
}}

.badge:hover {{
  transform: translateY(-2px);
  box-shadow: 0 8px 20px rgba(0, 0, 0, 0.08);
  border-color: #cbd5e1;
}}

.badge .icon {{
  font-size: 18px;
}}

/* ===== CARDS ===== */
.card {{
  background: #ffffff;
  border-radius: 24px;
  box-shadow: 0 10px 40px rgba(0, 0, 0, 0.08);
  padding: 40px;
  margin-bottom: 32px;
  transition: all 0.3s ease;
  border: 1px solid rgba(226, 232, 240, 0.6);
}}

.card:hover {{
  transform: translateY(-4px);
  box-shadow: 0 20px 60px rgba(0, 0, 0, 0.12);
}}

.card-header {{
  margin-bottom: 32px;
  padding-bottom: 20px;
  border-bottom: 2px solid #f1f5f9;
}}

.card h3 {{
  font-size: 24px;
  font-weight: 700;
  color: #334155;
  display: flex;
  align-items: center;
  gap: 12px;
}}

.card h3 .icon {{
  font-size: 28px;
}}

/* ===== STATS GRID ===== */
.stats-grid {{
  display: grid;
  grid-template-columns: repeat(auto-fit, minmax(260px, 1fr));
  gap: 24px;
}}

.stat-card {{
  text-align: center;
  padding: 32px 24px;
  background: linear-gradient(135deg, #f8fafc 0%, #f1f5f9 100%);
  border-radius: 20px;
  border: 2px solid #e2e8f0;
  transition: all 0.3s ease;
  position: relative;
  overflow: hidden;
}}

.stat-card::before {{
  content: '';
  position: absolute;
  top: 0;
  left: 0;
  right: 0;
  height: 4px;
  background: linear-gradient(90deg, #667eea 0%, #764ba2 100%);
  transform: scaleX(0);
  transition: transform 0.3s ease;
}}

.stat-card:hover {{
  transform: translateY(-6px);
  box-shadow: 0 12px 32px rgba(0, 0, 0, 0.1);
  border-color: #cbd5e1;
}}

.stat-card:hover::before {{
  transform: scaleX(1);
}}

.stat-icon {{
  font-size: 32px;
  margin-bottom: 12px;
  display: block;
}}

.stat-label {{
  font-size: 14px;
  font-weight: 600;
  color: #64748b;
  text-transform: uppercase;
  letter-spacing: 0.5px;
  margin-bottom: 8px;
}}

.stat-value {{
  font-weight: 800;
  font-size: 40px;
  background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
  -webkit-background-clip: text;
  -webkit-text-fill-color: transparent;
  background-clip: text;
  line-height: 1.2;
}}

.stat-card.overtime .stat-value {{
  background: linear-gradient(135deg, #ef4444 0%, #dc2626 100%);
  -webkit-background-clip: text;
  -webkit-text-fill-color: transparent;
  background-clip: text;
}}

/* ===== CHARTS GRID ===== */
.charts-grid {{
  display: grid;
  grid-template-columns: repeat(auto-fit, minmax(550px, 1fr));
  gap: 32px;
  margin-bottom: 32px;
}}

.chart-container {{
  position: relative;
  height: 400px;
}}

/* ===== WARNINGS ===== */
.warning-card {{
  background: linear-gradient(135deg, #fffbeb 0%, #fef3c7 100%);
  border: 2px solid #fbbf24;
}}

.warning-card .card-header {{
  border-bottom-color: #fbbf24;
}}

.warnings-container {{
  display: flex;
  flex-direction: column;
  gap: 12px;
}}

.warning-item {{
  display: flex;
  align-items: center;
  gap: 12px;
  padding: 16px 20px;
  background: rgba(255, 255, 255, 0.8);
  border-left: 4px solid #f59e0b;
  border-radius: 12px;
  color: #92400e;
  font-size: 15px;
  font-weight: 500;
  transition: all 0.2s ease;
}}

.warning-item:hover {{
  background: rgba(255, 255, 255, 1);
  transform: translateX(4px);
}}

.warning-icon {{
  font-size: 20px;
  flex-shrink: 0;
}}

/* ===== TABLES ===== */
.table-wrapper {{
  overflow-x: auto;
  margin-top: 20px;
  border-radius: 16px;
  border: 1px solid #e2e8f0;
}}

table {{
  width: 100%;
  border-collapse: separate;
  border-spacing: 0;
}}

thead {{
  background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
  position: sticky;
  top: 0;
  z-index: 10;
}}

thead th {{
  padding: 18px 20px;
  text-align: left;
  color: #ffffff;
  font-weight: 700;
  text-transform: uppercase;
  font-size: 13px;
  letter-spacing: 0.5px;
  white-space: nowrap;
}}

tbody tr {{
  transition: all 0.2s ease;
  border-bottom: 1px solid #f1f5f9;
}}

tbody tr:hover {{
  background: linear-gradient(135deg, #f8fafc 0%, #f1f5f9 100%);
  transform: scale(1.01);
}}

tbody tr:last-child {{
  border-bottom: none;
}}

tbody td {{
  padding: 16px 20px;
  font-size: 15px;
  color: #475569;
}}

tbody td:first-child {{
  font-weight: 600;
  color: #334155;
}}

/* ===== RESPONSIVE ===== */
@media (max-width: 768px) {{
  body {{
    padding: 20px 16px;
  }}

  .header {{
    padding: 32px 24px;
  }}

  .header-title {{
    font-size: 32px;
  }}

  .card {{
    padding: 24px;
  }}

  .stats-grid {{
    grid-template-columns: repeat(auto-fit, minmax(140px, 1fr));
    gap: 16px;
  }}

  .stat-card {{
    padding: 20px 16px;
  }}

  .stat-value {{
    font-size: 32px;
  }}

  .charts-grid {{
    grid-template-columns: 1fr;
  }}

  table {{
    font-size: 14px;
  }}

  thead th,
  tbody td {{
    padding: 12px 16px;
  }}
}}

/* ===== ANIMATIONS ===== */
@keyframes fadeInUp {{
  from {{
    opacity: 0;
    transform: translateY(20px);
  }}
  to {{
    opacity: 1;
    transform: translateY(0);
  }}
}}

.card {{
  animation: fadeInUp 0.6s ease-out;
}}

.card:nth-child(1) {{ animation-delay: 0.1s; }}
.card:nth-child(2) {{ animation-delay: 0.2s; }}
.card:nth-child(3) {{ animation-delay: 0.3s; }}
.card:nth-child(4) {{ animation-delay: 0.4s; }}
.card:nth-child(5) {{ animation-delay: 0.5s; }}
</style>
</head>
<body>
<div class="container">
  <!-- HEADER -->
  <div class="header">
    <h1 class="header-title">📊 Báo Cáo Tổng Hợp - Nhiều Nhân Viên</h1>
    <div class="header-meta">
      <span class="badge">
        <span class="icon">👥</span>
        <span>Số nhân viên: {stats['total_employees']}</span>
      </span>
      <span class="badge">
        <span class="icon">📅</span>
        <span>{stats['date_range']['start']} → {stats['date_range']['end']}</span>
      </span>
    </div>
  </div>

  <!-- STATS OVERVIEW -->
  <div class="card">
    <div class="stats-grid">
      <div class="stat-card">
        <span class="stat-icon">⏱️</span>
        <div class="stat-label">Tổng giờ</div>
        <div class="stat-value">{stats['total_hours']:.1f}h</div>
      </div>
      <div class="stat-card">
        <span class="stat-icon">📆</span>
        <div class="stat-label">Tổng ngày làm</div>
        <div class="stat-value">{stats['total_work_days']}</div>
      </div>
      <div class="stat-card">
        <span class="stat-icon">📊</span>
        <div class="stat-label">Trung bình/ngày</div>
        <div class="stat-value">{stats['avg_hours_per_day']:.1f}h</div>
      </div>
      <div class="stat-card overtime">
        <span class="stat-icon">🔥</span>
        <div class="stat-label">Tổng tăng ca</div>
        <div class="stat-value">{stats['overtime_info']['total_overtime']:.1f}h</div>
      </div>
    </div>
  </div>

  <!-- CHARTS -->
  <div class="charts-grid">
    <div class="card">
      <div class="card-header">
        <h3><span class="icon">👥</span>Phân bổ theo nhân viên</h3>
      </div>
      <div class="chart-container">
        <canvas id="employeeChart"></canvas>
      </div>
    </div>
    
    <div class="card">
      <div class="card-header">
        <h3><span class="icon">📈</span>Phân bổ theo dự án</h3>
      </div>
      <div class="chart-container">
        <canvas id="projectChart"></canvas>
      </div>
    </div>
  </div>

  <!-- WEEKLY CHART -->
  <div class="card">
    <div class="card-header">
      <h3><span class="icon">📊</span>Giờ làm theo tuần</h3>
    </div>
    <div class="chart-container">
      <canvas id="weeklyChart"></canvas>
    </div>
  </div>

  <!-- WARNINGS -->
  {warnings_html}

  <!-- EMPLOYEE DETAILS TABLE -->
  <div class="card">
    <div class="card-header">
      <h3><span class="icon">👤</span>Chi tiết theo nhân viên</h3>
    </div>
    <div class="table-wrapper">
      <table>
        <thead>
          <tr>
            <th>Nhân viên</th>
            <th>Tổng giờ</th>
            <th>Số ngày</th>
            <th>TB/ngày</th>
            <th>Tăng ca</th>
            <th>Phần trăm</th>
          </tr>
        </thead>
        <tbody>
          {''.join([
            f"<tr><td>{r['name']}</td><td>{r['total_hours']:.1f}h</td><td>{r['work_days']}</td><td>{r['avg_hours_per_day']:.1f}h</td><td>{r['total_overtime']:.1f}h</td><td>{r['percentage']:.1f}%</td></tr>"
            for _, r in emp.iterrows()
          ])}
        </tbody>
      </table>
    </div>
  </div>

  <!-- PROJECT DETAILS TABLE -->
  <div class="card">
    <div class="card-header">
      <h3><span class="icon">📋</span>Chi tiết theo dự án</h3>
    </div>
    <div class="table-wrapper">
      <table>
        <thead>
          <tr>
            <th>Dự án</th>
            <th>Tổng giờ</th>
            <th>Số ngày</th>
            <th>Số NV</th>
            <th>Phần trăm</th>
          </tr>
        </thead>
        <tbody>
          {''.join([
            f"<tr><td>{r['project']}</td><td>{r['total_hours']:.1f}h</td><td>{r['days_worked']}</td><td>{r['employees']}</td><td>{r['percentage']:.1f}%</td></tr>"
            for _, r in proj.iterrows()
          ])}
        </tbody>
      </table>
    </div>
  </div>
</div>

<script>
// Data
const empLabels = {json.dumps(emp_labels, ensure_ascii=False)};
const empHours = {json.dumps(emp_hours)};
const empColors = {json.dumps(emp_colors)};

const projLabels = {json.dumps(proj_labels, ensure_ascii=False)};
const projHours = {json.dumps(proj_hours)};
const projColors = {json.dumps(proj_colors)};

const weeklyLabels = {json.dumps(weekly_labels, ensure_ascii=False)};
const weeklyHours = {json.dumps(weekly_hours)};
const weeklyColors = {json.dumps(weekly_colors)};

// Chart.js Configuration
Chart.defaults.font.family = "'Inter', sans-serif";
Chart.defaults.font.size = 13;
Chart.defaults.color = '#475569';

// Employee Chart (Doughnut)
new Chart(document.getElementById('employeeChart').getContext('2d'), {{
  type: 'doughnut',
  data: {{
    labels: empLabels,
    datasets: [{{
      data: empHours,
      backgroundColor: empColors,
      borderWidth: 4,
      borderColor: '#ffffff',
      hoverOffset: 8
    }}]
  }},
  options: {{
    responsive: true,
    maintainAspectRatio: false,
    plugins: {{
      legend: {{
        position: 'bottom',
        labels: {{
          padding: 16,
          font: {{
            size: 13,
            weight: '600'
          }},
          usePointStyle: true,
          pointStyle: 'circle'
        }}
      }},
      tooltip: {{
        backgroundColor: 'rgba(0, 0, 0, 0.8)',
        padding: 12,
        titleFont: {{
          size: 14,
          weight: '700'
        }},
        bodyFont: {{
          size: 13
        }},
        cornerRadius: 8,
        callbacks: {{
          label: function(context) {{
            let label = context.label || '';
            let value = context.parsed || 0;
            let total = context.dataset.data.reduce((a, b) => a + b, 0);
            let percentage = ((value / total) * 100).toFixed(1);
            return label + ': ' + value.toFixed(1) + 'h (' + percentage + '%)';
          }}
        }}
      }}
    }}
  }}
}});

// Project Chart (Pie)
new Chart(document.getElementById('projectChart').getContext('2d'), {{
  type: 'pie',
  data: {{
    labels: projLabels,
    datasets: [{{
      data: projHours,
      backgroundColor: projColors,
      borderWidth: 4,
      borderColor: '#ffffff',
      hoverOffset: 8
    }}]
  }},
  options: {{
    responsive: true,
    maintainAspectRatio: false,
    plugins: {{
      legend: {{
        position: 'bottom',
        labels: {{
          padding: 16,
          font: {{
            size: 13,
            weight: '600'
          }},
          usePointStyle: true,
          pointStyle: 'circle'
        }}
      }},
      tooltip: {{
        backgroundColor: 'rgba(0, 0, 0, 0.8)',
        padding: 12,
        titleFont: {{
          size: 14,
          weight: '700'
        }},
        bodyFont: {{
          size: 13
        }},
        cornerRadius: 8,
        callbacks: {{
          label: function(context) {{
            let label = context.label || '';
            let value = context.parsed || 0;
            let total = context.dataset.data.reduce((a, b) => a + b, 0);
            let percentage = ((value / total) * 100).toFixed(1);
            return label + ': ' + value.toFixed(1) + 'h (' + percentage + '%)';
          }}
        }}
      }}
    }}
  }}
}});

// Weekly Chart (Bar)
new Chart(document.getElementById('weeklyChart').getContext('2d'), {{
  type: 'bar',
  data: {{
    labels: weeklyLabels,
    datasets: [{{
      label: 'Giờ làm việc',
      data: weeklyHours,
      backgroundColor: weeklyColors,
      borderRadius: 10,
      borderSkipped: false
    }}]
  }},
  options: {{
    responsive: true,
    maintainAspectRatio: false,
    plugins: {{
      legend: {{
        display: false
      }},
      tooltip: {{
        backgroundColor: 'rgba(0, 0, 0, 0.8)',
        padding: 12,
        titleFont: {{
          size: 14,
          weight: '700'
        }},
        bodyFont: {{
          size: 13
        }},
        cornerRadius: 8,
        callbacks: {{
          label: function(context) {{
            return 'Tổng giờ: ' + context.parsed.y.toFixed(1) + 'h';
          }}
        }}
      }}
    }},
    scales: {{
      y: {{
        beginAtZero: true,
        grid: {{
          color: 'rgba(226, 232, 240, 0.5)',
          drawBorder: false
        }},
        ticks: {{
          font: {{
            size: 12,
            weight: '600'
          }},
          color: '#64748b',
          callback: function(value) {{
            return value + 'h';
          }}
        }}
      }},
      x: {{
        grid: {{
          display: false,
          drawBorder: false
        }},
        ticks: {{
          font: {{
            size: 12,
            weight: '600'
          }},
          color: '#64748b'
        }}
      }}
    }}
  }}
}});
</script>
</body>
</html>"""

    Path(OUT_DASHBOARD_HTML).write_text(html, encoding="utf-8")
    print(f"✅ Đã tạo Dashboard tổng hợp: {OUT_DASHBOARD_HTML}")


def create_multi_user_excel(
    all_records: List[dict],
    stats: Dict[str, object],
    all_employees: Dict[str, List[dict]],
    OUT_EXCEL: str,
    STANDARD_HOURS: float = 8.0,
    all_dates: Optional[List[str]] = None,
) -> None:
    """Excel tổng hợp nhiều user (multi-sheet)."""
    try:
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        with pd.ExcelWriter(OUT_EXCEL, engine="openpyxl") as writer:
            header_fill = PatternFill(start_color="667EEA", end_color="667EEA", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            data_font = Font(name="Segoe UI", size=10, color="475569")
            center = Alignment(horizontal="center", vertical="center")
            thin = Border(
                left=Side(style="thin", color="E2E8F0"),
                right=Side(style="thin", color="E2E8F0"),
                top=Side(style="thin", color="E2E8F0"),
                bottom=Side(style="thin", color="E2E8F0"),
            )

            def fmt(ws):
                for c in range(1, ws.max_column + 1):
                    cell = ws.cell(row=1, column=c)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = center
                    cell.border = thin
                for r in range(2, ws.max_row + 1):
                    for c in range(1, ws.max_column + 1):
                        cell = ws.cell(row=r, column=c)
                        cell.font = data_font
                        cell.alignment = center
                        cell.border = thin
                for c in range(1, ws.max_column + 1):
                    ws.column_dimensions[get_column_letter(c)].width = 15

            # Tổng hợp
            df_summary = pd.DataFrame(
                {
                    "Chỉ số": ["Tổng nhân viên", "Tổng giờ", "Tổng ngày làm", "TB/Ngày", "Tổng tăng ca"],
                    "Giá trị": [
                        stats["total_employees"],
                        f"{stats['total_hours']:.1f}h",
                        stats["total_work_days"],
                        f"{stats['avg_hours_per_day']:.1f}h",
                        f"{stats['overtime_info']['total_overtime']:.1f}h",
                    ],
                }
            )
            df_summary.to_excel(writer, sheet_name="Tổng Hợp", index=False)
            fmt(writer.sheets["Tổng Hợp"])

            # Theo nhân viên
            df_emp = stats["employee_stats"].copy()
            df_emp.rename(
                columns={
                    "name": "Nhân viên",
                    "total_hours": "Tổng giờ",
                    "work_days": "Số ngày",
                    "avg_hours_per_day": "TB/Ngày",
                    "total_overtime": "Tăng ca",
                    "percentage": "%",
                },
                inplace=True,
            )
            df_emp.to_excel(writer, sheet_name="Theo Nhân Viên", index=False)
            fmt(writer.sheets["Theo Nhân Viên"])

            # Theo dự án
            df_proj = stats["project_stats"].copy()
            df_proj.rename(
                columns={
                    "project": "Dự án",
                    "total_hours": "Tổng giờ",
                    "days_worked": "Số ngày",
                    "employees": "Số NV",
                    "percentage": "%",
                },
                inplace=True,
            )
            df_proj.to_excel(writer, sheet_name="Theo Dự Án", index=False)
            fmt(writer.sheets["Theo Dự Án"])

            # Theo tuần
            df_week = stats["weekly_stats"].copy()
            df_week.rename(columns={"week": "Tuần", "total_hours": "Tổng giờ"}, inplace=True)
            df_week.to_excel(writer, sheet_name="Theo Tuần", index=False)
            fmt(writer.sheets["Theo Tuần"])

            # Dữ liệu gốc
            df_raw = pd.DataFrame(all_records)
            if not df_raw.empty and "project" in df_raw.columns:
                df_raw["project"] = df_raw["project"].apply(_canon_project_display)
            df_raw.rename(
                columns={"name": "Nhân viên", "date": "Ngày", "project": "Dự án", "hours": "Giờ"}, inplace=True
            )
            df_raw = df_raw.sort_values(["Nhân viên", "Ngày"])
            df_raw.to_excel(writer, sheet_name="Dữ Liệu Gốc", index=False)
            fmt(writer.sheets["Dữ Liệu Gốc"])

        print(f"✅ Đã tạo Excel tổng hợp: {OUT_EXCEL}")
    except Exception as e:
        print(f"⚠️ Lỗi tạo Excel: {e}")


def create_multi_user_pdf(stats: Dict[str, object], OUT_PDF: str) -> None:
    """PDF tổng hợp nhiều user."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        base_font = "Helvetica"
        bold_font = "Helvetica-Bold"
        fp = _pick_font_path()
        if fp:
            pdfmetrics.registerFont(TTFont("AppUnicode", fp))
            base_font = bold_font = "AppUnicode"

        doc = SimpleDocTemplate(OUT_PDF, pagesize=A4)
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "Title", parent=styles["Heading1"], fontName=bold_font, textColor=colors.HexColor("#667eea")
        )
        normal = ParagraphStyle("Normal", parent=styles["Normal"], fontName=base_font, fontSize=11)

        elements = []
        elements.append(Paragraph("BÁO CÁO TỔNG HỢP - NHIỀU NHÂN VIÊN", title_style))
        elements.append(Paragraph(f"Số nhân viên: {stats['total_employees']}", normal))
        elements.append(
            Paragraph(f"Khoảng: {stats['date_range']['start']} → {stats['date_range']['end']}", normal)
        )
        elements.append(Spacer(1, 0.25 * inch))

        data = [
            ["Chỉ Số", "Giá Trị"],
            ["Tổng nhân viên", str(stats["total_employees"])],
            ["Tổng giờ", f"{stats['total_hours']:.1f}h"],
            ["Tổng ngày làm", str(stats["total_work_days"])],
            ["TB/Ngày", f"{stats['avg_hours_per_day']:.1f}h"],
            ["Tổng tăng ca", f"{stats['overtime_info']['total_overtime']:.1f}h"],
        ]
        t = Table(data, colWidths=[3 * inch, 2.5 * inch])
        t.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), base_font),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#667eea")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                    ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ]
            )
        )
        elements.append(t)
        elements.append(Spacer(1, 0.2 * inch))

        elements.append(
            Paragraph("Phân tích theo nhân viên", ParagraphStyle("h2", parent=styles["Heading2"], fontName=bold_font))
        )
        emp = stats["employee_stats"].copy()
        emp["total_hours"] = emp["total_hours"].map(lambda x: f"{x:.1f}")
        emp["avg_hours_per_day"] = emp["avg_hours_per_day"].map(lambda x: f"{x:.1f}")
        emp["total_overtime"] = emp["total_overtime"].map(lambda x: f"{x:.1f}")
        emp["percentage"] = emp["percentage"].map(lambda x: f"{x:.1f}%")
        edata = [
            ["Nhân viên", "Tổng giờ", "Số ngày", "TB/Ngày", "Tăng ca", "%"]
        ] + emp[
            ["name", "total_hours", "work_days", "avg_hours_per_day", "total_overtime", "percentage"]
        ].values.tolist()
        te = Table(edata, colWidths=[1.8 * inch, 1 * inch, 0.8 * inch, 0.9 * inch, 0.9 * inch, 0.8 * inch])
        te.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), base_font),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#667eea")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                    ("ALIGN", (1, 1), (-1, -1), "CENTER"),
                ]
            )
        )
        elements.append(te)
        elements.append(Spacer(1, 0.25 * inch))
        elements.append(Paragraph(f"Tạo lúc: {datetime.now().strftime('%d/%m/%Y %H:%M')}", normal))

        doc.build(elements)
        print(f"✅ Đã tạo PDF tổng hợp: {OUT_PDF}")
    except Exception as e:
        print(f"⚠️ Lỗi tạo PDF: {e}")


# ======================================================================
#                              MAIN PIPELINE
# ======================================================================


def run_multi_user_pipeline(
    file_paths: List[str],
    out_dir: str,
    date_format: str = "day_first",
    standard_hours: float = 8.0,
) -> Dict[str, object]:
    """
    Pipeline chính (multi-user).
    """
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)

    OUT_MATRIX_XLSX = str(out / "project_matrix.xlsx")
    OUT_MULTI_XLSX = str(out / "work_report_complete.xlsx")
    OUT_DASHBOARD = str(out / "dashboard.html")
    OUT_PDF = str(out / "work_report.pdf")

    print("\n" + "=" * 60)
    print("🚀 BẮT ĐẦU XỬ LÝ MULTI-USER")
    print("=" * 60)

    print("\n📂 Đọc và xử lý các file CSV...")
    data = process_multiple_csvs(file_paths, date_format)
    if not data["all_records"]:
        raise RuntimeError("Không trích xuất được bản ghi nào từ các file CSV.")
    print(f"\n✅ Đã xử lý {len(data['all_employees'])} nhân viên, tổng {len(data['all_records'])} bản ghi")

    print("\n📊 Tính toán thống kê tổng hợp...")
    stats = calculate_multi_user_statistics(
        data["all_records"], data["all_employees"], standard_hours, standard_hours
    )

    print("\n📄 Tạo các file báo cáo...\n")
    print("1️⃣ Project Matrix Excel...")
    create_multi_user_matrix_excel(data["all_employees"], data["dataframes"], OUT_MATRIX_XLSX, date_format)

    print("\n2️⃣ Dashboard HTML...")
    create_multi_user_dashboard(stats, OUT_DASHBOARD, data["all_records"], data["dataframes"], date_format)

    print("\n3️⃣ Excel đa sheet...")
    create_multi_user_excel(
        data["all_records"], stats, data["all_employees"], OUT_MULTI_XLSX, standard_hours, data["all_dates"]
    )

    print("\n4️⃣ PDF Report...")
    create_multi_user_pdf(stats, OUT_PDF)

    print("\n" + "=" * 60)
    print("✅ HOÀN TẤT! Tất cả file đã được tạo thành công")
    print("=" * 60)

    return {
        "total_employees": len(data["all_employees"]),
        "employee_names": list(data["all_employees"].keys()),
        "total_records": len(data["all_records"]),
        "outputs": {"matrix": OUT_MATRIX_XLSX, "excel": OUT_MULTI_XLSX, "html": OUT_DASHBOARD, "pdf": OUT_PDF},
    }


if __name__ == "__main__":
    print("Đây là module core cho multi-user. Hãy chạy app_multi_user.py để mở GUI.")