# worklog_core.py
# -*- coding: utf-8 -*-

import re
from pathlib import Path
from datetime import datetime
from collections import defaultdict

import pandas as pd


# ============================ Helpers ============================

def _normalize_project_key(name: str) -> str:
    """
    Trả về 'khóa' ổn định cho tên dự án, chỉ xét chữ cái/số/Kanji/Kana.
    Loại bỏ khoảng trắng (kể cả full-width) và các dấu phân cách/phụ trợ.
    Ví dụ:
      '日本通運㈱神戸支店 / 東苅藻新倉庫計画' 
    -> '日本通運㈱神戸支店東苅藻新倉庫計画'
    """
    if not isinstance(name, str):
        name = '' if pd.isna(name) else str(name)
    s = name.strip()
    # Bỏ khoảng trắng (kể cả full-width) trước
    s = re.sub(r'[\s\u3000]+', '', s)
    # Bỏ các dấu câu/ngăn cách hay gặp (có thể mở rộng thêm khi cần)
    s = re.sub(r'[\/／・,，、\.\uFF0E\-\(\)（）\[\]【】「」『』:：;；…‧･]+', '', s)
    return s


def _normalize_decimal_commas(text: str) -> str:
    """8,5 / 8，5 / 8、5 -> 8.5 (chỉ khi giữa hai chữ số)."""
    s = str(text)
    return re.sub(r'(?<=\d)[,\uFF0C\u3001](?=\d)', '.', s).strip()


def _try_read_csv(path, encoding_list, seps, skiprows, DEBUG=False):
    """Thử nhiều encoding & separator cho tới khi đọc được."""
    last_err = None
    for enc in encoding_list:
        for sep in seps:
            try:
                df = pd.read_csv(path, encoding=enc, sep=sep, skiprows=skiprows, header=0)
                df.columns = [str(c).strip() for c in df.columns]

                time_cols = [col for col in df.columns if '時間' in col]
                if time_cols:
                    for col in time_cols:
                        df[col] = df[col].apply(_normalize_decimal_commas)

                df = df.dropna(axis=1, how='all')
                if len(df.columns) > 0:
                    if DEBUG and time_cols:
                        print("Mẫu '時間':", df[time_cols[0]].astype(str).head(5).tolist())
                    return df
            except Exception as e:
                last_err = e
                continue
    if last_err:
        raise last_err
    raise ValueError("Không đọc được CSV với các phương án đã thử.")


def _extract_employee_name_from_lines(lines, keywords=('氏名',)):
    """
    Tìm tên sau từ khóa (vd: '氏名') trong 30 dòng đầu.
    Hỗ trợ khoảng trắng full-width \u3000, dấu :/：, dấu phẩy ,/，/、, ngoặc kép.
    """
    kw_pat = r'(?:' + '|'.join(map(re.escape, keywords)) + r')'
    delims = r'[：:\s\u3000,，、]*'
    name_pat = r'"?(?P<name>[^,\s\u3000"“”]+)"?'
    rx = re.compile(kw_pat + delims + name_pat)

    for line in lines[:30]:
        if any(k in line for k in keywords):
            m = rx.search(line)
            if m:
                name = m.group('name').strip().strip('",，、')
                if name:
                    return name

            # Fallback: sau keyword tới dấu phẩy đầu
            for kw in keywords:
                if kw in line:
                    after = line.split(kw, 1)[1]
                    after = after.lstrip(' \t,，、:\u3000').strip()
                    after = after.split(',')[0].strip().strip('",，、')
                    if after:
                        return after

    # Fallback: lấy cụm JP dài nhất trong 10 dòng đầu
    jp_token_rx = re.compile(r'[ぁ-んァ-ン一-龥々ー]+')
    cand = ""
    for line in lines[:10]:
        for m in jp_token_rx.finditer(line):
            token = m.group(0).strip('",，、')
            if len(token) > len(cand):
                cand = token
    return cand if cand else "Unknown"


def _parse_date_safely(date_val, DATE_FORMAT_IN_CSV):
    """Parse ngày theo day_first hoặc month_first, fallback bằng pandas."""
    if pd.isna(date_val) or str(date_val).strip() == '':
        return None
    s = str(date_val).strip()

    # Phổ biến: YYYY/MM/DD
    if '/' in s:
        parts = s.split('/')
        if len(parts) == 3 and len(parts[0]) == 4:
            try:
                return datetime.strptime(s, '%Y/%m/%d')
            except ValueError:
                try:
                    y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
                    return datetime(y, m, d)
                except Exception:
                    pass

    if DATE_FORMAT_IN_CSV == 'month_first':
        fmts = ['%Y/%m/%d', '%m/%d/%Y', '%m/%d', '%Y-%m-%d']
    else:
        fmts = ['%d/%m/%Y', '%d/%m', '%Y/%m/%d', '%Y-%m-%d']

    for fmt in fmts:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass

    try:
        return pd.to_datetime(s, dayfirst=(DATE_FORMAT_IN_CSV == 'day_first'), errors='coerce')
    except Exception:
        return None


def _to_hours(val):
    """Chuyển nhiều định dạng giờ → float giờ."""
    if pd.isna(val):
        return 0.0
    s = str(val).strip()
    if s == '' or s.lower() == 'nan':
        return 0.0

    s = _normalize_decimal_commas(s)

    # hh:mm
    if ':' in s:
        parts = s.split(':')
        if len(parts) >= 2 and parts[0].strip().isdigit() and parts[1].strip().isdigit():
            return int(parts[0]) + int(parts[1]) / 60.0

    # X時間Y分 / XhYm / Xh / X
    m = re.match(r'^\s*(\d+(?:\.\d+)?)\s*(?:h|H|時間)?\s*(?:(\d{1,2})\s*(?:m|M|分))?\s*$', s)
    if m:
        h = float(m.group(1))
        mn = float(m.group(2)) if m.group(2) else 0.0
        return h + mn / 60.0

    try:
        return float(s)
    except Exception:
        return 0.0


# ============================ Parsing ============================

def parse_csv_data(file_path, DATE_FORMAT_IN_CSV='day_first'):
    """
    Đọc CSV 日報, tự tìm dòng header có '日付', trả về (df, employee_name).
    """
    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(str(file_path))

    # Đọc thô để dò header + quét tên
    try:
        raw = file_path.read_text(encoding='utf-8-sig')
    except UnicodeDecodeError:
        raw = file_path.read_text(encoding='cp932', errors='ignore')
    lines = raw.splitlines(True)

    # Tìm dòng header chứa '日付'
    header_line = -1
    for i, line in enumerate(lines[:200]):
        if '日付' in line:
            header_line = i
            break
    if header_line == -1:
        raise ValueError("Không tìm thấy header (含 '日付') trong file.")

    # Đọc DataFrame
    encodings = ['utf-8-sig', 'cp932', 'shift-jis']
    seps = [',', ';']
    df = _try_read_csv(str(file_path), encodings, seps, skiprows=header_line)

    # Lấy tên nhân viên
    employee_name = _extract_employee_name_from_lines(lines, keywords=('氏名',))
    if employee_name == "Unknown":
        m = re.search(r'\(([^)]+)\)', file_path.name)
        if m and m.group(1).strip():
            employee_name = m.group(1).strip()

    # Làm sạch
    df = df.dropna(how='all')
    if '日付' in df.columns:
        df = df[df['日付'].notna()]
        df = df[df['日付'].astype(str).str.contains(r'\d{4}[/-]\d{1,2}[/-]\d{1,2}', na=False)]

    return df, employee_name


# ===================== Trích xuất & Thống kê =====================

def extract_work_records(df, employee_name, DATE_FORMAT_IN_CSV='day_first'):
    """
    Trích xuất từ cột ngay sau '工事名' (cột giờ).
    Hỗ trợ nhiều cột '時間'.
    """
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]

    if '日付' not in df.columns or '工事名' not in df.columns:
        raise ValueError("Thiếu cột '日付' hoặc '工事名'.")

    proj_idx = list(df.columns).index('工事名')
    time_cols_idx = [i for i, col in enumerate(df.columns) if '時間' in col]
    hour_col_idx = None
    for idx in time_cols_idx:
        if idx > proj_idx:
            hour_col_idx = idx
            break
    if hour_col_idx is None:
        raise ValueError("Không tìm thấy cột giờ (含'時間') sau cột '工事名'.")

    hour_col = df.columns[hour_col_idx]

    records = []
    proj_canon = {}  # NEW: map key -> tên hiển thị chuẩn (tên đầu tiên gặp)

    def canon_display(name: str) -> str:  # NEW
        key = _normalize_project_key(name)
        if key not in proj_canon:
            proj_canon[key] = name.strip()
        return proj_canon[key]

    for _, row in df.iterrows():
        date_raw = row['日付']
        proj_raw = row['工事名']
        hours_raw = row[hour_col]

        if pd.isna(date_raw) or str(date_raw).strip() == '':
            continue

        d_obj = _parse_date_safely(date_raw, DATE_FORMAT_IN_CSV)
        if d_obj is None or pd.isna(d_obj):
            continue
        d_str = d_obj.strftime('%Y-%m-%d')

        project_lines = []
        if not pd.isna(proj_raw):
            project_lines = [p.strip() for p in str(proj_raw).split('\n') if str(p).strip() not in ('', 'nan', 'NaN')]

        hours_raw_str = '' if pd.isna(hours_raw) else str(hours_raw).strip()

        if '\n' in hours_raw_str:
            hours_list = [h.strip() for h in hours_raw_str.split('\n') if h.strip() not in ('', 'nan', 'NaN')]
            if len(project_lines) <= 1:
                total = sum(_to_hours(h) for h in hours_list)
                if total > 0 and project_lines:
                    proj_name = canon_display(project_lines[0])  # NEW
                    records.append({'name': employee_name, 'date': d_str, 'project': proj_name, 'hours': float(total)})
            else:
                for i, proj_name in enumerate(project_lines):
                    if i < len(hours_list):
                        hrs = _to_hours(hours_list[i])
                        if proj_name and hrs > 0:
                            proj_name = canon_display(proj_name)  # NEW
                            records.append({'name': employee_name, 'date': d_str, 'project': proj_name, 'hours': float(hrs)})
        else:
            total_hrs = _to_hours(hours_raw_str)
            if total_hrs > 0 and project_lines:
                merged = " / ".join(project_lines)
                merged = canon_display(merged)  # NEW: đảm bảo bản ghi gộp cũng dùng khóa chuẩn
                records.append({'name': employee_name, 'date': d_str, 'project': merged, 'hours': float(total_hrs)})

    return records

def analyze_overtime(records, STANDARD_HOURS=8.0, OVERTIME_THRESHOLD=8.0):
    """Tính tổng giờ/ngày và tạo cảnh báo tăng ca/thiếu giờ."""
    daily_hours = defaultdict(float)
    for r in records:
        daily_hours[r['date']] += float(r['hours'])

    overtime_records, warnings = [], []
    for d in sorted(daily_hours.keys()):
        total = daily_hours[d]
        if total > OVERTIME_THRESHOLD:
            ot = total - STANDARD_HOURS
            overtime_records.append({'date': d, 'total_hours': total, 'overtime': ot})
            warnings.append(f"⚠️  {d}: Làm {total:.1f}h (tăng ca {ot:.1f}h)")
        elif 0 < total < STANDARD_HOURS:
            warnings.append(f"ℹ️  {d}: Chỉ làm {total:.1f}h (thiếu {STANDARD_HOURS - total:.1f}h)")

    return {
        'overtime_records': overtime_records,
        'warnings': warnings,
        'total_overtime': sum(x['overtime'] for x in overtime_records)
    }


def calculate_statistics(records, df_original=None, STANDARD_HOURS=8.0, OVERTIME_THRESHOLD=8.0):
    """Sinh các bảng thống kê phục vụ Dashboard/Excel/PDF."""
    if not records:
        return {}

    df = pd.DataFrame(records)
    df['date_obj'] = pd.to_datetime(df['date'])

    total_hours = float(df['hours'].sum())
    work_days = int(df['date'].nunique())
    avg_hours_per_day = float(total_hours / work_days) if work_days else 0.0

    project_stats = df.groupby('project', as_index=False)['hours'] \
                      .agg(total_hours='sum', days_worked='count') \
                      .sort_values('total_hours', ascending=False)
    project_stats['percentage'] = (project_stats['total_hours'] / total_hours * 100).round(2) if total_hours > 0 else 0.0

    df['week'] = df['date_obj'].dt.isocalendar().week
    weekly_stats = df.groupby('week', as_index=False)['hours'].sum().rename(columns={'hours': 'total_hours'})

    overtime_info = analyze_overtime(records, STANDARD_HOURS, OVERTIME_THRESHOLD)

    return {
        'total_hours': total_hours,
        'work_days': work_days,
        'avg_hours_per_day': avg_hours_per_day,
        'project_stats': project_stats,
        'weekly_stats': weekly_stats,
        'overtime_info': overtime_info,
        'date_range': {
            'start': df['date_obj'].min().strftime('%Y-%m-%d'),
            'end': df['date_obj'].max().strftime('%Y-%m-%d')
        }
    }


# ============================ Outputs ============================

def _unique_sheet_name(original: str, used: set) -> str:
    """Tên sheet hợp lệ & duy nhất (<=31 ký tự)."""
    base = re.sub(r'[\\/*?:\[\]:]', '_', (original or '').strip()) or "Sheet"
    base = base[:28]
    cand, i = base, 1
    while cand in used:
        suffix = f"_{i}"
        cand = (base[:31 - len(suffix)]) + suffix
        i += 1
    used.add(cand)
    return cand


def create_project_matrix_excel(all_records_dict, df_original, OUT_MATRIX_XLSX, DATE_FORMAT_IN_CSV='day_first'):
    """
    Mỗi sheet = 1 công trình; Hàng = Tên người; Cột = ngày dd/mm theo CSV gốc.
    """
    if not all_records_dict:
        print("⚠️  Không có dữ liệu để tạo ma trận.")
        return

    all_records = []
    for _, recs in all_records_dict.items():
        all_records.extend(recs)
    if not all_records:
        print("⚠️  Không có bản ghi hợp lệ.")
        return

    # Danh sách ngày theo CSV gốc
    all_dates_dict = {}
    if df_original is not None and '日付' in df_original.columns:
        for date_val in df_original['日付']:
            if pd.isna(date_val) or str(date_val).strip() == '':
                continue
            d = _parse_date_safely(date_val, DATE_FORMAT_IN_CSV)
            if d is not None and not pd.isna(d):
                all_dates_dict[d.strftime('%d/%m')] = d
    all_dates_in_csv = [k for k, v in sorted(all_dates_dict.items(), key=lambda x: x[1])]

    df_all = pd.DataFrame(all_records)
    df_all['date_obj'] = pd.to_datetime(df_all['date'])
    df_all['date_formatted'] = df_all['date_obj'].dt.strftime('%d/%m')
    projects = list(df_all['project'].unique())

    used_names = set()
    try:
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        
        with pd.ExcelWriter(OUT_MATRIX_XLSX, engine='openpyxl') as writer:
            for project in projects:
                df_p = df_all[df_all['project'] == project].copy()
                pivot = df_p.pivot_table(
                    index='name',
                    columns='date_formatted',
                    values='hours',
                    aggfunc='sum',
                    fill_value=0
                )

                for d in all_dates_in_csv:
                    if d not in pivot.columns:
                        pivot[d] = None
                pivot = pivot.reindex(all_dates_in_csv, axis=1)

                pivot['Tổng'] = pivot.sum(axis=1, numeric_only=True)
                pivot.loc['TỔNG'] = pivot.sum(numeric_only=True)

                sheet_name = _unique_sheet_name(project, used_names)
                pivot.to_excel(writer, sheet_name=sheet_name)
                
                # === Formatting ===
                ws = writer.sheets[sheet_name]
                
                # Định nghĩa styles
                header_fill = PatternFill(start_color='667EEA', end_color='667EEA', fill_type='solid')
                header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
                total_fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
                total_font = Font(name='Segoe UI', size=10, bold=True, color='334155')
                data_font = Font(name='Segoe UI', size=10, color='475569')
                name_font = Font(name='Segoe UI', size=10, bold=True, color='334155')
                center_align = Alignment(horizontal='center', vertical='center')
                left_align = Alignment(horizontal='left', vertical='center')
                thin_border = Border(
                    left=Side(style='thin', color='E2E8F0'),
                    right=Side(style='thin', color='E2E8F0'),
                    top=Side(style='thin', color='E2E8F0'),
                    bottom=Side(style='thin', color='E2E8F0')
                )
                
                # Header row (row 1)
                for col in range(1, ws.max_column + 1):
                    cell = ws.cell(row=1, column=col)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = center_align
                    cell.border = thin_border
                
                # Data rows
                for row in range(2, ws.max_row + 1):
                    # Cột tên (A)
                    name_cell = ws.cell(row=row, column=1)
                    if name_cell.value == 'TỔNG':
                        name_cell.fill = total_fill
                        name_cell.font = total_font
                    else:
                        name_cell.font = name_font
                    name_cell.alignment = left_align
                    name_cell.border = thin_border
                    
                    # Các cột dữ liệu
                    for col in range(2, ws.max_column + 1):
                        cell = ws.cell(row=row, column=col)
                        
                        # Hàng TỔNG
                        if name_cell.value == 'TỔNG':
                            cell.fill = total_fill
                            cell.font = total_font
                        # Cột Tổng
                        elif ws.cell(row=1, column=col).value == 'Tổng':
                            cell.font = total_font
                        else:
                            cell.font = data_font
                        
                        cell.alignment = center_align
                        cell.border = thin_border
                        
                        # Format số
                        if isinstance(cell.value, (int, float)) and cell.value != 0:
                            cell.number_format = '0.0'
                
                # Auto-adjust column widths
                ws.column_dimensions['A'].width = 25  # Cột tên rộng hơn
                for col in range(2, ws.max_column + 1):
                    col_letter = get_column_letter(col)
                    if ws.cell(row=1, column=col).value == 'Tổng':
                        ws.column_dimensions[col_letter].width = 12
                    else:
                        ws.column_dimensions[col_letter].width = 10
                
                # Freeze panes (cố định header và cột tên)
                ws.freeze_panes = 'B2'
                
                print(f"  ✓ Sheet: {sheet_name}")

        print(f"✅ Đã tạo file ma trận công trình: {OUT_MATRIX_XLSX}")
    except ImportError:
        print("⚠️  Cần cài openpyxl: pip install openpyxl")
    except Exception as e:
        print(f"⚠️  Lỗi khi tạo file ma trận: {e}")

def _build_project_matrix_for_dashboard(records, df_original, DATE_FORMAT_IN_CSV='day_first'):
    """
    Tạo dữ liệu ma trận cho Dashboard (tab theo công trình).
    - dates: dd/mm theo thứ tự từ CSV gốc (cột 日付)
    - tables[project]: DataFrame đã format text để render (0/'' phù hợp)
    """
    # 1) Lấy danh sách ngày dd/mm theo CSV gốc
    all_dates_dict = {}
    if df_original is not None and '日付' in df_original.columns:
        for date_val in df_original['日付']:
            if pd.isna(date_val) or str(date_val).strip() == '':
                continue
            d = _parse_date_safely(date_val, DATE_FORMAT_IN_CSV)
            if d is not None and not pd.isna(d):
                all_dates_dict[d.strftime('%d/%m')] = d
    dates = [k for k, v in sorted(all_dates_dict.items(), key=lambda x: x[1])]

    # 2) Pivot theo từng project
    df_all = pd.DataFrame(records, columns=['name','date','project','hours']).copy()
    if df_all.empty:
        return [], dates, {}

    df_all['date_obj'] = pd.to_datetime(df_all['date'])
    df_all['date_col'] = df_all['date_obj'].dt.strftime('%d/%m')

    projects = list(df_all['project'].unique())
    tables = {}

    for project in projects:
        df_p = df_all[df_all['project'] == project].copy()
        pivot = df_p.pivot_table(
            index='name',
            columns='date_col',
            values='hours',
            aggfunc='sum',
            fill_value=0
        )

        # Bổ sung đủ ngày & theo thứ tự
        for d in dates:
            if d not in pivot.columns:
                pivot[d] = None
        pivot = pivot.reindex(dates, axis=1)

        # Hàng tổng & cột tổng
        # -> dùng to_numeric(sum).fillna(0) để loại nan
        total_cols = pivot.apply(pd.to_numeric, errors='coerce').sum(axis=0).fillna(0)
        pivot['Tổng'] = pivot.apply(pd.to_numeric, errors='coerce').sum(axis=1).fillna(0)
        pivot.loc['TỔNG'] = list(total_cols) + [float(total_cols.sum())]

        # ---- Format hiển thị ----
        #   - ô dữ liệu thường: 0 -> '' (trống), >0 -> "x.y"
        #   - riêng hàng TỔNG: mọi NaN/None -> "0", 0 -> "0"
        def _fmt_cell(x):
            # ô thường: để trống nếu không có số hoặc = 0
            if x is None or (isinstance(x, float) and pd.isna(x)):
                return ''
            try:
                xv = float(x)
                return '' if xv == 0.0 else f"{xv:.1f}"
            except Exception:
                return ''

        display = pivot.copy()
        # format toàn bảng trước
        display = display.applymap(_fmt_cell)
        # ép hàng TỔNG: rỗng -> "0"
        display.loc['TỔNG'] = display.loc['TỔNG'].replace('', '0')

        tables[project] = display

    return projects, dates, tables



def create_html_dashboard(stats, employee_name, OUT_DASHBOARD_HTML, records=None, df_original=None, DATE_FORMAT_IN_CSV='day_first'):
    """Tạo Dashboard HTML (Chart.js) + giao diện project matrix (tab/bảng)."""
    import json
    from pathlib import Path

    proj = stats['project_stats']
    week = stats['weekly_stats']
    ot   = stats['overtime_info']

    # ===== Biểu đồ (như cũ, có màu sắc)
    proj_labels = proj['project'].tolist()
    proj_hours  = proj['total_hours'].astype(float).tolist()
    proj_pct    = proj['percentage'].astype(float).tolist() if len(proj) else []
    weekly_labels = [f"Tuần {int(w)}" for w in week['week'].tolist()]
    weekly_hours  = week['total_hours'].astype(float).tolist()

    palette = [
        "#6366F1","#22C55E","#F59E0B","#EF4444","#14B8A6","#A855F7",
        "#3B82F6","#84CC16","#F97316","#06B6D4","#EC4899","#10B981",
        "#8B5CF6","#EAB308","#FB7185","#0EA5E9","#34D399","#F43F5E"
    ]
    proj_colors   = [palette[i % len(palette)] for i in range(len(proj_labels))]
    weekly_colors = [palette[i % len(palette)] for i in range(len(weekly_labels))]

    # ===== Dữ liệu ma trận để gắn vào dashboard
    projects, dates_ddmm, table_dict = _build_project_matrix_for_dashboard(
        records or [], df_original, DATE_FORMAT_IN_CSV
    )

    # Render HTML cho các tab + bảng
    tabs_html = ''.join([
        f'<button class="tab-btn" data-target="tb_{i}">{p}</button>'
        for i, p in enumerate(projects)
    ]) or '<em>Không có dữ liệu ma trận.</em>'

    def render_table_html(df):
        # header
        ths = ''.join([f'<th>{d}</th>' for d in df.columns])
        # rows
        rows = []
        for idx, row in df.iterrows():
            tds = ''.join([f'<td>{row[c]}</td>' for c in df.columns])
            rows.append(f"<tr><th>{idx}</th>{tds}</tr>")
        return f"""
        <div class="table-wrap">
          <table class="matrix">
            <thead><tr><th>Nhân viên</th>{ths}</tr></thead>
            <tbody>{''.join(rows)}</tbody>
          </table>
        </div>
        """

    tables_html = ''.join([
        f'<div class="tab-panel" id="tb_{i}" style="display:{ "block" if i==0 else "none" }">{render_table_html(table_dict[p])}</div>'
        for i, p in enumerate(projects)
    ])

    html = f"""<!DOCTYPE html>
<html lang="vi"><head>
<meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Báo Cáo Làm Việc - {employee_name}</title>
<script src="https://cdnjs.cloudflare.com/ajax/libs/Chart.js/3.9.1/chart.min.js"></script>
<style>
* {{ margin: 0; padding: 0; box-sizing: border-box; }}

body {{
    font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
    background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
    padding: 32px 20px;
    min-height: 100vh;
    color: #1e293b;
}}

.container {{
    max-width: 1400px;
    margin: 0 auto;
}}

/* Header */
.header {{
    background: #ffffff;
    border-radius: 20px;
    box-shadow: 0 10px 40px rgba(0, 0, 0, 0.1);
    padding: 32px 40px;
    margin-bottom: 32px;
    border: 1px solid rgba(255, 255, 255, 0.2);
}}

.header h1 {{
    font-size: 32px;
    font-weight: 700;
    background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
    margin-bottom: 16px;
    display: flex;
    align-items: center;
    gap: 12px;
}}

.badge {{
    display: inline-block;
    padding: 10px 20px;
    border-radius: 12px;
    background: linear-gradient(135deg, #f8fafc 0%, #e2e8f0 100%);
    color: #475569;
    font-weight: 600;
    font-size: 14px;
    margin-right: 12px;
    border: 1px solid #e2e8f0;
    transition: all 0.3s ease;
}}

.badge:hover {{
    transform: translateY(-2px);
    box-shadow: 0 4px 12px rgba(0, 0, 0, 0.1);
}}

/* Cards */
.card, .panel {{
    background: #ffffff;
    border-radius: 20px;
    box-shadow: 0 10px 40px rgba(0, 0, 0, 0.08);
    padding: 32px;
    margin-bottom: 32px;
    border: 1px solid rgba(255, 255, 255, 0.2);
    transition: all 0.3s ease;
}}

.card:hover {{
    transform: translateY(-4px);
    box-shadow: 0 20px 60px rgba(0, 0, 0, 0.12);
}}

.card h3, .panel h3 {{
    font-size: 20px;
    font-weight: 700;
    color: #334155;
    margin-bottom: 24px;
    display: flex;
    align-items: center;
    gap: 10px;
}}

/* Stats Grid */
.grid {{
    display: grid;
    grid-template-columns: repeat(auto-fit, minmax(260px, 1fr));
    gap: 24px;
}}

.stat {{
    text-align: center;
    padding: 24px;
    background: linear-gradient(135deg, #f8fafc 0%, #f1f5f9 100%);
    border-radius: 16px;
    border: 2px solid #e2e8f0;
    transition: all 0.3s ease;
}}

.stat:hover {{
    transform: translateY(-4px);
    border-color: #667eea;
    box-shadow: 0 8px 24px rgba(102, 126, 234, 0.2);
}}

.stat > div:first-child {{
    font-size: 14px;
    font-weight: 600;
    color: #64748b;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    margin-bottom: 12px;
}}

.stat .n {{
    font-weight: 800;
    font-size: 36px;
    background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
    margin-top: 8px;
}}

/* Warnings */
.warn, .info {{
    padding: 16px 20px;
    border-radius: 12px;
    margin: 12px 0;
    font-size: 14px;
    line-height: 1.6;
    border-left: 4px solid;
    transition: all 0.3s ease;
}}

.warn {{
    background: linear-gradient(135deg, #fffbeb 0%, #fef3c7 100%);
    border-color: #f59e0b;
    color: #92400e;
}}

.info {{
    background: linear-gradient(135deg, #ecfeff 0%, #cffafe 100%);
    border-color: #06b6d4;
    color: #164e63;
}}

.warn:hover, .info:hover {{
    transform: translateX(4px);
    box-shadow: 0 4px 12px rgba(0, 0, 0, 0.1);
}}

/* Tables */
.table-wrap {{
    overflow: auto;
    max-width: 100%;
    border-radius: 12px;
    box-shadow: inset 0 2px 8px rgba(0, 0, 0, 0.05);
}}

table {{
    width: 100%;
    border-collapse: separate;
    border-spacing: 0;
    margin-top: 16px;
}}

th, td {{
    padding: 14px 16px;
    text-align: left;
    white-space: nowrap;
    font-size: 14px;
}}

thead th {{
    background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
    color: #ffffff;
    font-weight: 700;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    font-size: 12px;
    position: sticky;
    top: 0;
    z-index: 10;
}}

tbody tr {{
    transition: all 0.2s ease;
    border-bottom: 1px solid #f1f5f9;
}}

tbody tr:hover {{
    background: linear-gradient(135deg, #f8fafc 0%, #f1f5f9 100%);
    transform: scale(1.01);
}}

tbody td {{
    color: #475569;
}}

.matrix thead th:first-child, .matrix tbody th {{
    position: sticky;
    left: 0;
    background: #f8fafc;
    z-index: 5;
    font-weight: 700;
    color: #334155;
    border-right: 2px solid #e2e8f0;
}}

.matrix tbody th {{
    background: linear-gradient(135deg, #f8fafc 0%, #f1f5f9 100%);
}}

/* Tabs */
.tabs {{
    display: flex;
    gap: 12px;
    flex-wrap: wrap;
    margin-bottom: 24px;
    padding-bottom: 16px;
    border-bottom: 2px solid #e2e8f0;
}}

.tab-btn {{
    background: linear-gradient(135deg, #f8fafc 0%, #f1f5f9 100%);
    color: #475569;
    border: 2px solid #e2e8f0;
    border-radius: 12px;
    padding: 12px 24px;
    cursor: pointer;
    font-weight: 600;
    font-size: 14px;
    transition: all 0.3s ease;
    position: relative;
    overflow: hidden;
}}

.tab-btn::before {{
    content: '';
    position: absolute;
    top: 0;
    left: -100%;
    width: 100%;
    height: 100%;
    background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
    transition: left 0.3s ease;
    z-index: -1;
}}

.tab-btn:hover {{
    transform: translateY(-2px);
    box-shadow: 0 4px 12px rgba(102, 126, 234, 0.3);
    border-color: #667eea;
}}

.tab-btn.active {{
    background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
    color: #ffffff;
    border-color: #667eea;
    box-shadow: 0 4px 16px rgba(102, 126, 234, 0.4);
}}

.tab-btn.active::before {{
    left: 0;
}}

.section-title {{
    margin: 0 0 16px;
    display: flex;
    align-items: center;
    gap: 12px;
    font-size: 22px;
}}

.section-title span {{
    font-size: 24px;
}}

/* Chart containers */
canvas {{
    max-height: 350px;
}}

/* Animations */
@keyframes fadeIn {{
    from {{
        opacity: 0;
        transform: translateY(20px);
    }}
    to {{
        opacity: 1;
        transform: translateY(0);
    }}
}}

.card, .panel, .header {{
    animation: fadeIn 0.6s ease-out;
}}

/* Responsive */
@media (max-width: 768px) {{
    body {{
        padding: 16px 12px;
    }}
    
    .header, .card, .panel {{
        padding: 20px;
        border-radius: 16px;
    }}
    
    .header h1 {{
        font-size: 24px;
    }}
    
    .badge {{
        font-size: 12px;
        padding: 8px 16px;
    }}
    
    .stat .n {{
        font-size: 28px;
    }}
    
    .tabs {{
        gap: 8px;
    }}
    
    .tab-btn {{
        padding: 10px 16px;
        font-size: 13px;
    }}
}}

/* Scrollbar styling */
.table-wrap::-webkit-scrollbar {{
    height: 8px;
    width: 8px;
}}

.table-wrap::-webkit-scrollbar-track {{
    background: #f1f5f9;
    border-radius: 4px;
}}

.table-wrap::-webkit-scrollbar-thumb {{
    background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
    border-radius: 4px;
}}

.table-wrap::-webkit-scrollbar-thumb:hover {{
    background: linear-gradient(135deg, #5a67d8 0%, #6b3fa0 100%);
}}
</style>
</head><body>
<div class="container">
  <div class="header">
    <h1>📊 Báo Cáo Làm Việc</h1>
    <div>
      <span class="badge">👤 Nhân viên: {employee_name}</span>
      <span class="badge">📅 Từ {stats['date_range']['start']} đến {stats['date_range']['end']}</span>
    </div>
  </div>

  <div class="card"><div class="grid">
    <div class="stat"><div>⏱️ Tổng giờ</div><div class="n">{stats['total_hours']:.1f}h</div></div>
    <div class="stat"><div>📆 Số ngày</div><div class="n">{stats['work_days']}</div></div>
    <div class="stat"><div>📊 TB/ngày</div><div class="n">{stats['avg_hours_per_day']:.1f}h</div></div>
    <div class="stat"><div>🔥 Tăng ca</div><div class="n" style="background: linear-gradient(135deg, #ef4444 0%, #dc2626 100%); -webkit-background-clip: text; -webkit-text-fill-color: transparent;">{stats['overtime_info']['total_overtime']:.1f}h</div></div>
  </div></div>

  {('<div class="panel"><h3>⚠️ Cảnh báo & Lưu ý</h3>' + ''.join([f'<div class="{{"warn" if "tăng ca" in w else "info"}}">{w}</div>' for w in ot['warnings']]) + '</div>') if ot['warnings'] else ''}

  <div class="grid">
    <div class="card"><h3>📈 Phân bổ theo dự án</h3><canvas id="project"></canvas></div>
    <div class="card"><h3>📊 Giờ theo tuần</h3><canvas id="weekly"></canvas></div>
  </div>

  <div class="panel">
    <h3 class="section-title">🧩 <span>Ma trận theo công trình</span></h3>
    <div class="tabs">{tabs_html}</div>
    {tables_html}
  </div>

  <div class="panel">
    <h3>📋 Chi tiết theo dự án</h3>
    <table>
      <thead><tr><th>Dự án</th><th>Tổng giờ</th><th>Số ngày</th><th>%</th></tr></thead>
      <tbody>
        {''.join([
          f"<tr><td>{r['project']}</td><td>{float(r['total_hours']):.1f}h</td><td>{int(r['days_worked'])}</td><td>{float(r['percentage']):.1f}%</td></tr>"
          for _, r in proj.iterrows()
        ])}
      </tbody>
    </table>
  </div>
</div>

<script>
const projLabels   = {json.dumps(proj_labels, ensure_ascii=False)};
const projHours    = {json.dumps(proj_hours)};
const projPct      = {json.dumps(proj_pct)};
const projColors   = {json.dumps(proj_colors)};
const weeklyLabels = {json.dumps(weekly_labels, ensure_ascii=False)};
const weeklyHours  = {json.dumps(weekly_hours)};
const weeklyColors = {json.dumps(weekly_colors)};

new Chart(document.getElementById('project').getContext('2d'), {{
  type: 'pie',
  data: {{
    labels: projLabels,
    datasets: [{{
      data: projHours,
      backgroundColor: projColors,
      borderColor: '#ffffff',
      borderWidth: 3,
      hoverOffset: 10
    }}]
  }},
  options: {{
    responsive: true,
    maintainAspectRatio: true,
    plugins: {{
      legend: {{ 
        position: 'bottom',
        labels: {{
          padding: 15,
          font: {{ size: 13, weight: '600' }},
          usePointStyle: true,
          pointStyle: 'circle'
        }}
      }},
      tooltip: {{ 
        backgroundColor: 'rgba(0, 0, 0, 0.8)',
        padding: 12,
        cornerRadius: 8,
        titleFont: {{ size: 14, weight: 'bold' }},
        bodyFont: {{ size: 13 }},
        callbacks: {{
          label: (ctx) => (ctx.label || '') + ': ' + (ctx.parsed || 0).toFixed(1) + 'h (' + (projPct[ctx.dataIndex] || 0) + '%)'
        }}
      }}
    }}
  }}
}});

new Chart(document.getElementById('weekly').getContext('2d'), {{
  type: 'bar',
  data: {{
    labels: weeklyLabels,
    datasets: [{{
      label: 'Giờ làm việc',
      data: weeklyHours,
      backgroundColor: weeklyColors,
      borderColor: weeklyColors,
      borderWidth: 0,
      borderRadius: 8,
      hoverBackgroundColor: weeklyColors.map(c => c + 'dd')
    }}]
  }},
  options: {{
    responsive: true,
    maintainAspectRatio: true,
    plugins: {{ 
      legend: {{ display: false }},
      tooltip: {{
        backgroundColor: 'rgba(0, 0, 0, 0.8)',
        padding: 12,
        cornerRadius: 8,
        titleFont: {{ size: 14, weight: 'bold' }},
        bodyFont: {{ size: 13 }}
      }}
    }},
    scales: {{ 
      y: {{ 
        beginAtZero: true,
        grid: {{ color: '#f1f5f9' }},
        ticks: {{ font: {{ size: 12 }} }}
      }},
      x: {{
        grid: {{ display: false }},
        ticks: {{ font: {{ size: 12 }} }}
      }}
    }}
  }}
}});

// Tabs cho ma trận
document.querySelectorAll('.tab-btn').forEach((btn, i) => {{
  btn.addEventListener('click', () => {{
    document.querySelectorAll('.tab-btn').forEach(b => b.classList.remove('active'));
    document.querySelectorAll('.tab-panel').forEach(p => p.style.display = 'none');
    btn.classList.add('active');
    document.getElementById(btn.dataset.target).style.display = 'block';
  }});
  if (i === 0) btn.classList.add('active');
}});
</script>
</body></html>"""
    Path(OUT_DASHBOARD_HTML).write_text(html, encoding='utf-8')
    print(f"✅ Đã tạo Dashboard HTML: {OUT_DASHBOARD_HTML}")



def _pick_font_path():
    """Chọn font Unicode cho PDF (ưu tiên font kèm theo app nếu có)."""
    import sys
    from pathlib import Path as P
    meipass = getattr(sys, '_MEIPASS', None)
    base_dir = P(meipass if meipass else P(__file__).parent)
    candidates = [
        base_dir / "fonts" / "NotoSans-Regular.ttf",
        base_dir / "fonts" / "NotoSansJP-Regular.ttf",
        P(r"C:\Windows\Fonts\segoeui.ttf"),
        P(r"C:\Windows\Fonts\arialuni.ttf"),
        P(r"C:\Windows\Fonts\DejaVuSans.ttf"),
    ]
    for p in candidates:
        if p.exists():
            return str(p)
    return None


def create_pdf_report(stats, employee_name, OUT_PDF_REPORT):
    """Tạo báo cáo PDF (ReportLab) với font Unicode."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        base_font = 'Helvetica'
        bold_font = 'Helvetica-Bold'
        fp = _pick_font_path()
        if fp:
            pdfmetrics.registerFont(TTFont('AppUnicode', fp))
            base_font = bold_font = 'AppUnicode'

        doc = SimpleDocTemplate(OUT_PDF_REPORT, pagesize=A4)
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontName=bold_font, textColor=colors.HexColor('#667eea'))
        normal = ParagraphStyle('Normal', parent=styles['Normal'], fontName=base_font, fontSize=11, leading=14)

        elements = []
        elements.append(Paragraph("BÁO CÁO LÀM VIỆC", title_style))
        elements.append(Paragraph(f"Nhân viên: {employee_name}", normal))
        elements.append(Paragraph(f"Khoảng thời gian: {stats['date_range']['start']} → {stats['date_range']['end']}", normal))
        elements.append(Spacer(1, 0.25*inch))

        # Tổng quan
        data = [['Chỉ Số','Giá Trị'],
                ['Tổng giờ', f"{stats['total_hours']:.1f}h"],
                ['Số ngày', str(stats['work_days'])],
                ['TB/ngày', f"{stats['avg_hours_per_day']:.1f}h"],
                ['Tổng tăng ca', f"{stats['overtime_info']['total_overtime']:.1f}h"]]
        t = Table(data, colWidths=[3*inch, 2.5*inch])
        t.setStyle(TableStyle([
            ('FONTNAME', (0,0), (-1,-1), base_font),
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#667eea')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('GRID', (0,0), (-1,-1), 0.3, colors.grey),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ]))
        elements.append(t); elements.append(Spacer(1, 0.2*inch))

        # Dự án
        proj = stats['project_stats'].copy()
        proj['total_hours'] = proj['total_hours'].map(lambda x: f"{float(x):.1f}")
        proj['percentage']  = proj['percentage'].map(lambda x: f"{float(x):.1f}%")
        pdata = [['Dự Án','Tổng Giờ','Số Ngày','%']] + proj[['project','total_hours','days_worked','percentage']].values.tolist()
        from reportlab.platypus import Table as T2, TableStyle as TS2  # tránh shadow
        tp = T2(pdata, colWidths=[3*inch, 1*inch, 1*inch, 1.2*inch])
        tp.setStyle(TS2([
            ('FONTNAME', (0,0), (-1,-1), base_font),
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#667eea')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('GRID', (0,0), (-1,-1), 0.3, colors.grey),
            ('ALIGN', (1,1), (-1,-1), 'CENTER'),
        ]))
        elements.append(Paragraph("Phân tích theo dự án", ParagraphStyle('h2', parent=styles['Heading2'], fontName=bold_font)))
        elements.append(tp); elements.append(Spacer(1, 0.2*inch))

        # Overtime
        if stats['overtime_info']['overtime_records']:
            ot = pd.DataFrame(stats['overtime_info']['overtime_records'])
            ot['total_hours'] = ot['total_hours'].map(lambda x: f"{float(x):.1f}")
            ot['overtime']    = ot['overtime'].map(lambda x: f"{float(x):.1f}")
            odata = [['Ngày','Tổng Giờ','Tăng Ca']] + ot[['date','total_hours','overtime']].values.tolist()
            to = T2(odata, colWidths=[2.2*inch, 1.5*inch, 1.5*inch])
            to.setStyle(TS2([
                ('FONTNAME', (0,0), (-1,-1), base_font),
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#667eea')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.white),
                ('GRID', (0,0), (-1,-1), 0.3, colors.grey),
                ('ALIGN', (1,1), (-1,-1), 'CENTER'),
            ]))
            elements.append(Paragraph("Cảnh báo tăng ca", ParagraphStyle('h2', parent=styles['Heading2'], fontName=bold_font)))
            elements.append(to)
        else:
            elements.append(Paragraph("Không có ngày tăng ca vượt ngưỡng.", normal))

        elements.append(Spacer(1, 0.25*inch))
        elements.append(Paragraph(f"Tạo lúc: {datetime.now().strftime('%d/%m/%Y %H:%M')}", normal))
        doc.build(elements)
        print(f"✅ Đã tạo PDF: {OUT_PDF_REPORT}")
    except ImportError:
        print("⚠️  Cần cài reportlab: pip install reportlab")
    except Exception as e:
        print(f"⚠️  Lỗi tạo PDF: {e}")


def create_multi_sheet_excel(
    records,
    stats,
    OUT_MULTI_XLSX,
    STANDARD_HOURS=8.0,
    all_dates_from_csv=None
):
    """
    Excel nhiều sheet: Tổng Hợp, Theo Ngày, Theo Dự Án, Theo Tuần, Tăng Ca, Dữ Liệu Gốc.
    """
    if not stats:
        print("⚠️  Không có dữ liệu để tạo Excel.")
        return
    try:
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        
        with pd.ExcelWriter(OUT_MULTI_XLSX, engine='openpyxl') as writer:
            # Định nghĩa styles chung
            header_fill = PatternFill(start_color='667EEA', end_color='667EEA', fill_type='solid')
            header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
            label_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
            label_font = Font(name='Segoe UI', size=10, bold=True, color='334155')
            data_font = Font(name='Segoe UI', size=10, color='475569')
            center_align = Alignment(horizontal='center', vertical='center')
            left_align = Alignment(horizontal='left', vertical='center')
            thin_border = Border(
                left=Side(style='thin', color='E2E8F0'),
                right=Side(style='thin', color='E2E8F0'),
                top=Side(style='thin', color='E2E8F0'),
                bottom=Side(style='thin', color='E2E8F0')
            )
            
            def format_sheet(ws, has_header=True, label_col=None):
                """Áp dụng format chung cho sheet"""
                if has_header:
                    # Header row
                    for col in range(1, ws.max_column + 1):
                        cell = ws.cell(row=1, column=col)
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = center_align
                        cell.border = thin_border
                
                # Data rows
                for row in range(2 if has_header else 1, ws.max_row + 1):
                    for col in range(1, ws.max_column + 1):
                        cell = ws.cell(row=row, column=col)
                        
                        # Cột label (nếu có)
                        if label_col and col == label_col:
                            cell.fill = label_fill
                            cell.font = label_font
                            cell.alignment = left_align
                        else:
                            cell.font = data_font
                            cell.alignment = center_align if col > (label_col or 0) else left_align
                        
                        cell.border = thin_border
                        
                        # Format số
                        if isinstance(cell.value, (int, float)):
                            if 'giờ' in str(ws.cell(row=1, column=col).value).lower() or 'hours' in str(ws.cell(row=1, column=col).value).lower():
                                cell.number_format = '0.0'
                
                # Auto-adjust columns
                for col in range(1, ws.max_column + 1):
                    col_letter = get_column_letter(col)
                    max_length = 0
                    for row in range(1, min(ws.max_row + 1, 100)):  # Sample first 100 rows
                        cell_value = str(ws.cell(row=row, column=col).value or '')
                        max_length = max(max_length, len(cell_value))
                    ws.column_dimensions[col_letter].width = min(max(max_length + 2, 12), 50)
                
                # Freeze panes
                if has_header:
                    ws.freeze_panes = 'A2'
            
            # ===== Sheet 1: Tổng Hợp =====
            df_summary = pd.DataFrame({
                'Chỉ Số': ['Tổng giờ','Số ngày','TB/Ngày','Tổng tăng ca','Bắt đầu','Kết thúc'],
                'Giá Trị': [
                    f"{stats['total_hours']:.1f}h",
                    stats['work_days'],
                    f"{stats['avg_hours_per_day']:.1f}h",
                    f"{stats['overtime_info']['total_overtime']:.1f}h",
                    stats['date_range']['start'],
                    stats['date_range']['end']
                ]
            })
            df_summary.to_excel(writer, sheet_name='Tổng Hợp', index=False)
            format_sheet(writer.sheets['Tổng Hợp'], has_header=True, label_col=1)

            # ===== Sheet 2: Theo Ngày =====
            df_rec = pd.DataFrame(records) if records else pd.DataFrame(columns=['date','project','hours'])
            daily = df_rec.groupby('date', as_index=False) \
                          .agg(hours=('hours','sum'),
                               projects=('project', lambda x: ', '.join(pd.unique(x))))
            daily['Tăng Ca'] = daily['hours'].apply(lambda x: f"{x - STANDARD_HOURS:.1f}h" if x > STANDARD_HOURS else '-')
            daily.rename(columns={'date':'Ngày','hours':'Tổng Giờ','projects':'Dự Án'}, inplace=True)
            daily = daily.sort_values('Ngày')
            daily.to_excel(writer, sheet_name='Theo Ngày', index=False)
            format_sheet(writer.sheets['Theo Ngày'], has_header=True)

            # ===== Sheet 3: Theo Dự Án =====
            df_project = stats['project_stats'].copy()
            df_project.rename(columns={
                'project': 'Dự Án',
                'total_hours': 'Tổng Giờ',
                'days_worked': 'Số Ngày',
                'percentage': 'Phần Trăm (%)'
            }, inplace=True)
            df_project.to_excel(writer, sheet_name='Theo Dự Án', index=False)
            format_sheet(writer.sheets['Theo Dự Án'], has_header=True)

            # ===== Sheet 4: Theo Tuần =====
            df_weekly = stats['weekly_stats'].copy()
            df_weekly.rename(columns={
                'week': 'Tuần',
                'total_hours': 'Tổng Giờ'
            }, inplace=True)
            df_weekly.to_excel(writer, sheet_name='Theo Tuần', index=False)
            format_sheet(writer.sheets['Theo Tuần'], has_header=True)

            # ===== Sheet 5: Tăng Ca =====
            if stats['overtime_info']['overtime_records']:
                ot = pd.DataFrame(stats['overtime_info']['overtime_records'])
                ot.rename(columns={'date':'Ngày','total_hours':'Tổng Giờ','overtime':'Giờ Tăng Ca'}, inplace=True)
                ot.to_excel(writer, sheet_name='Tăng Ca', index=False)
                
                ws_ot = writer.sheets['Tăng Ca']
                format_sheet(ws_ot, has_header=True)
                
                # Highlight overtime cells
                orange_fill = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
                for row in range(2, ws_ot.max_row + 1):
                    for col in range(2, ws_ot.max_column + 1):
                        cell = ws_ot.cell(row=row, column=col)
                        if col == ws_ot.max_column:  # Cột Giờ Tăng Ca
                            cell.fill = orange_fill

            # ===== Sheet 6: Dữ Liệu Gốc =====
            base_cols = ['date','project','hours']
            df_src = df_rec[base_cols].copy() if not df_rec.empty else pd.DataFrame(columns=base_cols)

            if all_dates_from_csv:
                have = set(df_src['date'].tolist())
                missing_rows = [{'date': d, 'project': '', 'hours': None}
                                for d in all_dates_from_csv if d not in have]
                if missing_rows:
                    df_src = pd.concat([df_src, pd.DataFrame(missing_rows)], ignore_index=True)

            df_src = df_src.sort_values('date', kind='stable')
            df_src.rename(columns={
                'date': 'Ngày',
                'project': 'Dự Án',
                'hours': 'Giờ'
            }, inplace=True)
            df_src.to_excel(writer, sheet_name='Dữ Liệu Gốc', index=False)
            format_sheet(writer.sheets['Dữ Liệu Gốc'], has_header=True)

        print(f"✅ Đã tạo Excel đa sheet: {OUT_MULTI_XLSX}")
    except ImportError:
        print("⚠️  Cần cài openpyxl: pip install openpyxl")
    except Exception as e:
        print(f"⚠️  Lỗi tạo Excel: {e}")



# ============================ API cho GUI ============================

def run_pipeline(file_path, out_dir, date_format='day_first', standard_hours=8.0):
    out_dir = Path(out_dir); out_dir.mkdir(parents=True, exist_ok=True)
    OUT_MATRIX_XLSX = str(out_dir / "project_matrix.xlsx")
    OUT_MULTI_XLSX  = str(out_dir / "work_report_complete.xlsx")
    OUT_DASHBOARD   = str(out_dir / "dashboard.html")
    OUT_PDF         = str(out_dir / "work_report.pdf")

    df, employee_name = parse_csv_data(file_path, DATE_FORMAT_IN_CSV=date_format)
    records = extract_work_records(df, employee_name, DATE_FORMAT_IN_CSV=date_format)
    if not records:
        raise RuntimeError("Không trích xuất được bản ghi nào từ CSV.")

    print("📊 Tạo project_matrix.xlsx (giữ nguyên)...")
    create_project_matrix_excel({employee_name: records}, df, OUT_MATRIX_XLSX, DATE_FORMAT_IN_CSV=date_format)

    print("📐 Tính thống kê & tăng ca...")
    stats = calculate_statistics(records, df_original=df, STANDARD_HOURS=standard_hours, OVERTIME_THRESHOLD=standard_hours)

    # 👇 Danh sách tất cả ngày xuất hiện trong CSV gốc (duy nhất, có thứ tự)
    all_dates = []
    if '日付' in df.columns:
        seen = set()
        for val in df['日付']:
            if pd.isna(val) or str(val).strip() == '':
                continue
            d = _parse_date_safely(val, date_format)
            if d is not None and not pd.isna(d):
                key = d.strftime('%Y-%m-%d')
                if key not in seen:
                    seen.add(key)
                    all_dates.append(key)

    print("🌐 Dashboard HTML...")
    create_html_dashboard(stats, employee_name, OUT_DASHBOARD,records=records,df_original=df,DATE_FORMAT_IN_CSV=date_format)

    print("📚 Excel multi-sheet...")
    # 👇 truyền all_dates vào để sheet 'Dữ Liệu Gốc' có đủ ngày
    create_multi_sheet_excel(records, stats, OUT_MULTI_XLSX, STANDARD_HOURS=standard_hours, all_dates_from_csv=all_dates)

    print("🖨️ PDF...")
    create_pdf_report(stats, employee_name, OUT_PDF)

    return {
        "employee_name": employee_name,
        "outputs": {
            "matrix": OUT_MATRIX_XLSX,
            "excel": OUT_MULTI_XLSX,
            "html": OUT_DASHBOARD,
            "pdf": OUT_PDF
        }
    }



if __name__ == "__main__":
    print("worklog_core là module lõi. Hãy chạy app.py để mở ứng dụng GUI.")
